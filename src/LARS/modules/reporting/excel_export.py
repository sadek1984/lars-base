"""
Excel export of an answer's table (lars_service /api/lars/query).

    write_export(question, df, source, info_rows) -> filename in export_dir()
    answer_info(engine, question, text, source)  -> rows for the "معلومات" sheet
    export_path(filename)                         -> Path, or None if not a safe export name

Workbook (right-to-left): "النتائج" = the table (bold header, fitted widths,
the same display labels as the chat answer); "معلومات" = question, source,
the semantic interpretation line, filters, period, generated_at, row count.
Read-only: nothing here changes routing or answers.
"""
from __future__ import annotations

import os
import re
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from modules.query.messages import EU_MRL_COLUMN_LABELS

REPO = Path(__file__).resolve().parents[4]
_SAFE_NAME = re.compile(r"[\w\-]+\.xlsx")
_MONTHS_AR = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو",
              "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
MAX_WIDTH = 60


def export_dir() -> Path:
    return Path(os.environ.get("LARS_EXPORT_DIR", REPO / "exports"))


def valid_export_name(filename: str) -> bool:
    """A plain *.xlsx name: no "..", no slashes, word characters and "-" only."""
    return bool(filename) and ".." not in filename and "/" not in filename \
        and "\\" not in filename and bool(_SAFE_NAME.fullmatch(filename))


def export_path(filename: str) -> Optional[Path]:
    """The file for `filename` if it is a valid export name that exists in
    export_dir(); None otherwise."""
    if not valid_export_name(filename):
        return None
    base = export_dir().resolve()
    path = (base / filename).resolve()
    if path.parent != base or not path.is_file():
        return None
    return path


def _slug(question: str, limit: int = 40) -> str:
    words = re.findall(r"[^\W_]+", question or "")
    return "_".join(words)[:limit].strip("_") or "export"


# ── "معلومات" ────────────────────────────────────────────────────────────────

def _join(values) -> str:
    return "، ".join(str(v) for v in values if v)


def _handler_filters(ctx: Optional[dict]) -> Tuple[str, str]:
    if not ctx:
        return "—", "—"
    res = ctx.get("resolution")
    parts = []
    for label, values in (("التصنيف", [ctx.get("category_key")]),
                          ("المنتج", ctx.get("detected_samples") or []),
                          ("البلدية", getattr(res, "municipalities", None) or []),
                          ("الحي", ctx.get("detected_neighborhoods") or []),
                          ("المبيد", [ctx.get("detected_pesticide")])):
        if _join(values):
            parts.append(f"{label}: {_join(values)}")
    period = ctx.get("detected_period_label")
    fragment = ctx.get("detected_period") or ""
    month = re.search(r"date_part\('month', .*\) = (\d+)$", fragment)
    if not period and month:
        period = f"شهر {_MONTHS_AR[int(month.group(1)) - 1]}"
    return (" | ".join(parts) or "بدون تصفية"), (period or ("—" if fragment else "كامل البيانات"))


def _semantic_filters(ans) -> Tuple[str, str]:
    from modules.semantic.answer import CATEGORY_AR
    r = ans.resolved
    parts = []
    for label, values in (("التصنيف", [CATEGORY_AR.get(c, c) for c in r.categories]),
                          ("المنتج", r.products), ("البلدية", r.municipalities),
                          ("الحي", r.neighborhoods), ("المبيد", r.pesticides)):
        if _join(values):
            parts.append(f"{label}: {_join(values)}")
    if r.spec.scope.value == "non_compliant":
        parts.append("العينات غير المطابقة فقط")
    period = f"{r.period[0]} إلى {r.period[1]}" if r.period else "كامل البيانات"
    return (" | ".join(parts) or "بدون تصفية"), period


def answer_info(engine, question: str, text: str, source: str) -> List[Tuple[str, str]]:
    """Rows for the "معلومات" sheet, from what the engine recorded for the
    last answer (engine may be None, e.g. the voice priority shortcut)."""
    ans = getattr(engine, "last_semantic", None) if source == "semantic" else None
    if ans is not None and getattr(ans, "resolved", None) is not None:
        filters, period = _semantic_filters(ans)
        interpretation = (text or "").splitlines()[0] if text else ""
    else:
        filters, period = _handler_filters(getattr(engine, "last_context", None))
        interpretation = ""
    rows = [("السؤال", question), ("المصدر", source)]
    if source == "semantic":
        rows.append(("تفسير السؤال", interpretation))
    rows += [("التصفية", filters), ("الفترة", period)]
    return rows


# ── workbook ─────────────────────────────────────────────────────────────────

def _fit(ws) -> None:
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, MAX_WIDTH)


def write_export(question: str, df: pd.DataFrame, source: str,
                 info_rows: List[Tuple[str, str]], now: Optional[datetime] = None) -> str:
    """Write the workbook; returns its file name (in export_dir())."""
    now = now or datetime.now()
    out_dir = export_dir()
    out_dir.mkdir(parents=True, exist_ok=True)
    name = f"{now:%Y%m%d-%H%M%S}-{uuid.uuid4().hex[:6]}_{_slug(question)}.xlsx"

    table = df.rename(columns=EU_MRL_COLUMN_LABELS)
    table = table.astype(object).where(table.notna(), None)
    wb = Workbook()
    ws = wb.active
    ws.title = "النتائج"
    ws.sheet_view.rightToLeft = True
    ws.append([str(c) for c in table.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in table.itertuples(index=False):
        ws.append([v.item() if hasattr(v, "item") else v for v in row])
    ws.freeze_panes = "A2"
    _fit(ws)

    info = wb.create_sheet("معلومات")
    info.sheet_view.rightToLeft = True
    rows = list(info_rows) + [("تاريخ الإنشاء", now.isoformat(timespec="seconds")),
                              ("عدد الصفوف", len(table))]
    for key, value in rows:
        info.append([key, value])
    for cell in info["A"]:
        cell.font = Font(bold=True)
    for row in info.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _fit(info)
    wb.save(out_dir / name)
    return name
