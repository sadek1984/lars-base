"""
Excel export in lars_service: /api/lars/query adds "file_url" for answers
with a table; GET /api/lars/export/{filename} serves only plain *.xlsx names
from the export directory.

Run:  pytest tests/test_excel_export.py -v
"""
import logging
import sys
import time
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
LARS = REPO / "src" / "LARS"
DB_PATH = LARS / "data" / "lars_data_demo.duckdb"
sys.path.insert(0, str(LARS))
sys.path.insert(0, str(REPO))

pytestmark = pytest.mark.skipif(not DB_PATH.exists(), reason="lars_data_demo.duckdb not present")
pytest.importorskip("fastapi")
pytest.importorskip("httpx")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from modules.semantic import fallback as fb  # noqa: E402
from modules.semantic.llm import SpecProvider  # noqa: E402

NUTS_NC = "كم عدد العينات غير المطابقة في المكسرات؟"


class Nuts(SpecProvider):
    name = "fake"

    def fill(self, question, system_prompt):
        return '{"metric": "noncompliant_count", "filters": {"category": ["المكسرات"]}}'


@pytest.fixture(scope="module")
def engine():
    logging.disable(logging.CRITICAL)
    from modules.query.core_query_engine import CoreQueryEngine
    e = CoreQueryEngine(db_path=str(DB_PATH))
    e.semantic_mode, e._semantic = "off", None
    return e


@pytest.fixture
def client(engine, tmp_path, monkeypatch):
    import lars_service
    monkeypatch.setenv("LARS_EXPORT_DIR", str(tmp_path / "exports"))
    monkeypatch.setattr(lars_service, "_engine", engine)
    return TestClient(lars_service.app)


@pytest.fixture(scope="module")
def bank():
    import csv
    with (REPO / "phase0" / "questions_bank.csv").open(encoding="utf-8") as f:
        return {r["id"]: r["question"] for r in csv.DictReader(f)}


def ask(client, q):
    r = client.post("/api/lars/query", json={"query": q})
    assert r.status_code == 200
    return r.json()


def download(client, out, tmp_path):
    g = client.get(out["file_url"])
    assert g.status_code == 200
    assert g.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    cd = g.headers["content-disposition"]
    assert cd.startswith('attachment; filename="') and ".xlsx\"; filename*=UTF-8''" in cd
    path = tmp_path / "dl.xlsx"
    path.write_bytes(g.content)
    return load_workbook(path)


def info(wb):
    return {k: v for k, v in wb["معلومات"].iter_rows(values_only=True)}


@pytest.mark.parametrize("qid", ["D002", "D011"])
def test_handler_answer_exports_its_table(client, engine, bank, tmp_path, qid):
    out = ask(client, bank[qid])
    text, df = engine.process(bank[qid])[:2]
    assert out["answer"] == text and out["source"] == "handler"          # the answer is unchanged
    assert out["file_url"].startswith("/api/lars/export/") and out["file_url"].endswith(".xlsx")
    wb = download(client, out, tmp_path)
    assert wb.sheetnames == ["النتائج", "معلومات"]
    res = wb["النتائج"]
    assert res.sheet_view.rightToLeft and wb["معلومات"].sheet_view.rightToLeft
    assert all(c.font.b for c in res[1]) and res.max_row == len(df) + 1
    assert res.column_dimensions["A"].width >= len(str(res["A1"].value))
    meta = info(wb)
    assert meta["السؤال"] == bank[qid] and meta["المصدر"] == "handler" and meta["عدد الصفوف"] == len(df)
    assert "تفسير السؤال" not in meta and meta["الفترة"] == "كامل البيانات" and meta["تاريخ الإنشاء"]


def test_d011_uses_the_display_labels(client, bank, tmp_path):
    wb = download(client, ask(client, bank["D011"]), tmp_path)
    header = [c.value for c in wb["النتائج"][1]]
    assert "violations" not in header and "تجاوزت الحد الأقصى الأوروبي" in header


def test_semantic_answer_exports_with_its_interpretation(client, engine, tmp_path):
    engine.semantic_mode = "live"
    engine._semantic = fb.SemanticFallback(str(DB_PATH), "live", provider=Nuts(), log_path=tmp_path / "log.jsonl")
    try:
        out = ask(client, NUTS_NC)
    finally:
        engine.semantic_mode, engine._semantic = "off", None
    assert out["source"] == "semantic" and out["refusal"] is None
    wb = download(client, out, tmp_path)
    rows = list(wb["النتائج"].iter_rows(values_only=True))
    assert rows[0] == ("samples", "non_compliant") and rows[1][1] == 5
    meta = info(wb)
    assert meta["المصدر"] == "semantic" and meta["تفسير السؤال"] == out["answer"].splitlines()[0]
    assert meta["التصفية"] == "التصنيف: المكسرات" and meta["الفترة"] == "كامل البيانات"


def test_period_is_recorded(client, tmp_path):
    wb = download(client, ask(client, "ما هي العينات غير المطابقة في شهر مارس؟"), tmp_path)
    assert info(wb)["الفترة"] == "شهر مارس"


def test_refusal_has_no_file(client, tmp_path):
    out = ask(client, NUTS_NC)                                           # off: category refusal
    assert out["refusal"] == "category" and "file_url" not in out
    assert not (tmp_path / "exports").exists() or not any((tmp_path / "exports").iterdir())


def test_answer_without_table_has_no_file(client):
    out = ask(client, "ما سعر الطماطم؟")                                 # not understood, no table
    assert "file_url" not in out


def test_export_failure_keeps_the_answer(client, monkeypatch, bank):
    import modules.reporting.excel_export as xe
    monkeypatch.setattr(xe, "write_export", lambda *a, **k: (_ for _ in ()).throw(OSError("disk full")))
    out = ask(client, bank["D002"])
    assert out["success"] and "file_url" not in out


@pytest.mark.parametrize("name, status", [
    ("x.csv", 400), ("a..b.xlsx", 400), ("..\\x.xlsx", 400), ("x.xlsx.exe", 400),
    ("nothere.xlsx", 404), ("..%2F..%2Fetc%2Fpasswd", 404),
])
def test_export_rejects_unsafe_names(client, name, status):
    assert client.get(f"/api/lars/export/{name}").status_code == status


def test_export_path_stays_in_the_directory(tmp_path, monkeypatch):
    from modules.reporting.excel_export import export_path
    monkeypatch.setenv("LARS_EXPORT_DIR", str(tmp_path / "exports"))
    (tmp_path / "exports").mkdir()
    (tmp_path / "secret.xlsx").write_bytes(b"x")
    (tmp_path / "exports" / "ok.xlsx").write_bytes(b"x")
    assert export_path("ok.xlsx") == (tmp_path / "exports" / "ok.xlsx").resolve()
    assert export_path("../secret.xlsx") is None and export_path("secret.xlsx") is None


def test_file_names(client, bank):
    a, b = ask(client, bank["D002"]), ask(client, bank["D002"])
    assert a["file_url"] != b["file_url"]                                # unique per request
    name = a["file_url"].rsplit("/", 1)[1]
    assert name[:8] == time.strftime("%Y%m%d")
